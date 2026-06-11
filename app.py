import os
import json
import datetime
import requests
import shutil
import re
from typing import List, Optional, Dict, Any
from fastapi import FastAPI, HTTPException, UploadFile, File, Form
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel

# Import document parsers
from pypdf import PdfReader
import docx
import openpyxl

app = FastAPI(title="247 Operations Shift Handover Agent API")

# Setup directories
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DATA_DIR = os.path.join(BASE_DIR, "data")
TEMP_DIR = os.path.join(BASE_DIR, "temp_uploads")
FRONTEND_DIR = os.path.join(BASE_DIR, "frontend")

os.makedirs(DATA_DIR, exist_ok=True)
os.makedirs(TEMP_DIR, exist_ok=True)
os.makedirs(FRONTEND_DIR, exist_ok=True)

# Helper paths
LOGS_FILE = os.path.join(DATA_DIR, "logs.json")
SOPS_FILE = os.path.join(DATA_DIR, "sops.json")
CONFIG_FILE = os.path.join(DATA_DIR, "config.json")

# Ensure files exist
if not os.path.exists(LOGS_FILE):
    with open(LOGS_FILE, "w", encoding="utf-8") as f:
        json.dump([], f)
if not os.path.exists(SOPS_FILE):
    with open(SOPS_FILE, "w", encoding="utf-8") as f:
        json.dump([], f)
if not os.path.exists(CONFIG_FILE):
    with open(CONFIG_FILE, "w", encoding="utf-8") as f:
        json.dump({
            "api_key": "",
            "model_name": "minimax/minimax-m2.5",
            "imap_server": "imap.gmail.com",
            "imap_user": "",
            "imap_pass": "",
            "imap_enabled": False
        }, f, indent=2)

# File DB functions
def read_json_file(file_path: str) -> List[Dict[str, Any]]:
    try:
        with open(file_path, "r", encoding="utf-8") as f:
            return json.load(f)
    except Exception:
        return []

def write_json_file(file_path: str, data: Any):
    with open(file_path, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)

def read_config() -> Dict[str, Any]:
    # Read environment variable first
    env_key = os.environ.get("MAAS_API_KEY")
    config = {
        "api_key": "",
        "model_name": "minimax/minimax-m2.5",
        "imap_server": "imap.gmail.com",
        "imap_user": "",
        "imap_pass": "",
        "imap_enabled": False
    }
    try:
        if os.path.exists(CONFIG_FILE):
            with open(CONFIG_FILE, "r", encoding="utf-8") as f:
                config.update(json.load(f))
    except Exception:
        pass
    if env_key:
        config["api_key"] = env_key
    return config

def write_config(config: Dict[str, Any]):
    write_json_file(CONFIG_FILE, config)

# Models
class LogItem(BaseModel):
    id: Optional[str] = None
    type: str  # incident, maintenance, note
    component: str
    severity: str  # low, medium, high
    status: str  # resolving, resolved, pending
    content: str
    created_at: Optional[str] = None
    resolved_at: Optional[str] = None

class ChatMessage(BaseModel):
    role: str
    content: str

class ChatRequest(BaseModel):
    message: str
    history: List[ChatMessage]

# Document parsers
def parse_txt(file_path: str) -> str:
    for encoding in ['utf-8', 'utf-16', 'latin-1', 'cp1252']:
        try:
            with open(file_path, "r", encoding=encoding) as f:
                return f.read()
        except UnicodeDecodeError:
            continue
    raise ValueError("Could not decode file with common encodings.")

def parse_pdf(file_path: str) -> str:
    reader = PdfReader(file_path)
    text_parts = []
    for i, page in enumerate(reader.pages):
        text = page.extract_text()
        if text:
            text_parts.append(text)
    return "\n".join(text_parts)

def parse_docx(file_path: str) -> str:
    doc = docx.Document(file_path)
    text_parts = []
    for p in doc.paragraphs:
        if p.text.strip():
            text_parts.append(p.text)
    return "\n".join(text_parts)

def parse_xlsx(file_path: str) -> str:
    wb = openpyxl.load_workbook(file_path, data_only=True)
    try:
        text_parts = []
        for sheet in wb.worksheets:
            text_parts.append(f"--- Sheet: {sheet.title} ---")
            for row in sheet.iter_rows(values_only=True):
                row_str = " | ".join([str(val) for val in row if val is not None])
                if row_str.strip():
                    text_parts.append(row_str)
        return "\n".join(text_parts)
    finally:
        wb.close()

# APIs
@app.get("/api/logs")
def get_logs():
    return read_json_file(LOGS_FILE)

@app.post("/api/logs")
def create_log(item: LogItem):
    logs = read_json_file(LOGS_FILE)
    item.id = f"log-{int(datetime.datetime.now().timestamp() * 1000)}"
    if not item.created_at:
        item.created_at = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    if item.status == "resolved" and not item.resolved_at:
        item.resolved_at = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    
    logs.append(item.dict())
    write_json_file(LOGS_FILE, logs)
    return item

@app.put("/api/logs/{log_id}")
def update_log(log_id: str, item: LogItem):
    logs = read_json_file(LOGS_FILE)
    for index, log in enumerate(logs):
        if log["id"] == log_id:
            # Check resolved timestamp transition
            if item.status == "resolved" and log["status"] != "resolved" and not item.resolved_at:
                item.resolved_at = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            elif item.status != "resolved":
                item.resolved_at = None
                
            item.id = log_id
            if not item.created_at:
                item.created_at = log["created_at"]
            
            logs[index] = item.dict()
            write_json_file(LOGS_FILE, logs)
            return item
    raise HTTPException(status_code=404, detail="Log not found")

@app.delete("/api/logs/{log_id}")
def delete_log(log_id: str):
    logs = read_json_file(LOGS_FILE)
    filtered_logs = [log for log in logs if log["id"] != log_id]
    if len(filtered_logs) == len(logs):
        raise HTTPException(status_code=404, detail="Log not found")
    write_json_file(LOGS_FILE, filtered_logs)
    return {"status": "success"}

@app.get("/api/sops")
def get_sops():
    return read_json_file(SOPS_FILE)

@app.delete("/api/sops/{sop_id}")
def delete_sop(sop_id: str):
    sops = read_json_file(SOPS_FILE)
    filtered_sops = [sop for sop in sops if sop["id"] != sop_id]
    if len(filtered_sops) == len(sops):
        raise HTTPException(status_code=404, detail="SOP not found")
    write_json_file(SOPS_FILE, filtered_sops)
    return {"status": "success"}

@app.post("/api/sops/upload")
async def upload_sop(file: UploadFile = File(...), title: str = Form(...)):
    ext = os.path.splitext(file.filename)[1].lower()
    temp_file_path = os.path.join(TEMP_DIR, f"temp_{file.filename}")
    
    with open(temp_file_path, "wb") as buffer:
        shutil.copyfileobj(file.file, buffer)
        
    try:
        content = ""
        if ext == ".txt":
            content = parse_txt(temp_file_path)
        elif ext == ".pdf":
            content = parse_pdf(temp_file_path)
        elif ext in [".docx", ".doc"]:
            content = parse_docx(temp_file_path)
        elif ext in [".xlsx", ".xls"]:
            content = parse_xlsx(temp_file_path)
        else:
            raise HTTPException(status_code=400, detail=f"Unsupported file format: {ext}")
        
        # Save to SOP Database
        sops = read_json_file(SOPS_FILE)
        sop_item = {
            "id": f"sop-{int(datetime.datetime.now().timestamp() * 1000)}",
            "title": title or file.filename,
            "filename": file.filename,
            "content": content
        }
        sops.append(sop_item)
        write_json_file(SOPS_FILE, sops)
        
        return {"status": "success", "data": sop_item}
        
    except Exception as e:
        import traceback
        tb = traceback.format_exc()
        with open(os.path.join(DATA_DIR, "error.log"), "w", encoding="utf-8") as f_err:
            f_err.write(tb)
        raise HTTPException(status_code=500, detail=f"Failed to parse file: {str(e)}. Traceback: {tb}")
    finally:
        # Cleanup temp file
        try:
            if os.path.exists(temp_file_path):
                os.remove(temp_file_path)
        except Exception as cleanup_err:
            print(f"Non-fatal error cleaning up temp file: {cleanup_err}")

@app.get("/api/config")
def get_sys_config():
    config = read_config()
    key = config.get("api_key", "")
    masked_key = ""
    if key:
        if len(key) > 10:
            masked_key = key[:6] + "•" * (len(key) - 10) + key[-4:]
        else:
            masked_key = "•" * len(key)
            
    imap_pass = config.get("imap_pass", "")
    masked_imap_pass = ""
    if imap_pass:
        masked_imap_pass = "•" * len(imap_pass)
        
    return {
        "api_key": masked_key,
        "has_key": bool(key),
        "model_name": config.get("model_name", "minimax/minimax-m2.5"),
        "imap_server": config.get("imap_server", "imap.gmail.com"),
        "imap_user": config.get("imap_user", ""),
        "imap_pass": masked_imap_pass,
        "imap_enabled": config.get("imap_enabled", False)
    }

@app.post("/api/config")
def update_sys_config(config_data: Dict[str, Any]):
    new_key = config_data.get("api_key", "").strip()
    
    # Read config from file only to update it, without env var overlay,
    # to avoid saving env vars back into config.json.
    local_config = {
        "api_key": "",
        "model_name": "minimax/minimax-m2.5",
        "imap_server": "imap.gmail.com",
        "imap_user": "",
        "imap_pass": "",
        "imap_enabled": False
    }
    try:
        if os.path.exists(CONFIG_FILE):
            with open(CONFIG_FILE, "r", encoding="utf-8") as f:
                local_config.update(json.load(f))
    except Exception:
        pass
        
    # If the key is masked (contains bullet points or stars), ignore updating the actual key
    if new_key and "•" not in new_key and "*" not in new_key:
        local_config["api_key"] = new_key
    elif not new_key:
        local_config["api_key"] = ""
        
    if "model_name" in config_data:
        local_config["model_name"] = config_data["model_name"]
        
    if "imap_server" in config_data:
        local_config["imap_server"] = config_data["imap_server"]
    if "imap_user" in config_data:
        local_config["imap_user"] = config_data["imap_user"]
    if "imap_pass" in config_data:
        new_pass = config_data["imap_pass"]
        if new_pass and "•" not in new_pass and "*" not in new_pass:
            local_config["imap_pass"] = new_pass
        elif not new_pass:
            local_config["imap_pass"] = ""
    if "imap_enabled" in config_data:
        local_config["imap_enabled"] = bool(config_data["imap_enabled"])
        
    write_config(local_config)
    return {"status": "success"}

@app.get("/api/handover")
def generate_handover(sender: Optional[str] = None, receiver: Optional[str] = None):
    logs = read_json_file(LOGS_FILE)
    
    incidents_resolving = [l for l in logs if l["type"] == "incident" and l["status"] != "resolved"]
    incidents_resolved = [l for l in logs if l["type"] == "incident" and l["status"] == "resolved"]
    maintenances = [l for l in logs if l["type"] == "maintenance"]
    notes = [l for l in logs if l["type"] == "note"]
    
    now_str = datetime.datetime.now().strftime("%d/%m/%Y %H:%M")
    
    report = []
    report.append(f"# BÁO CÁO BÀN GIAO CA TRỰC VẬN HÀNH (247 OPERATIONS)")
    report.append(f"*Thời gian lập báo cáo: {now_str}*")
    report.append(f"\n---")
    
    # 1. Sự cố đang xử lý
    report.append(f"\n## 1. Sự cố đang xử lý ({len(incidents_resolving)})")
    if not incidents_resolving:
        report.append("- Không có sự cố tồn đọng đang xử lý.")
    else:
        for idx, log in enumerate(incidents_resolving, 1):
            severity_tag = f"[{log['severity'].upper()}]"
            status_tag = f"[{log['status'].upper()}]"
            report.append(f"{idx}. **{log['component']}** {severity_tag} {status_tag}")
            report.append(f"   - **Thời gian phát hiện:** {log['created_at']}")
            report.append(f"   - **Chi tiết:** {log['content']}")
            
    # 2. Sự cố đã khắc phục trong ca
    report.append(f"\n## 2. Sự cố đã khắc phục trong ca ({len(incidents_resolved)})")
    if not incidents_resolved:
        report.append("- Không có sự cố phát sinh và khắc phục trong ca.")
    else:
        for idx, log in enumerate(incidents_resolved, 1):
            report.append(f"{idx}. **{log['component']}** [RESOLVED]")
            report.append(f"   - **Thời gian:** {log['created_at']} -> {log['resolved_at']}")
            report.append(f"   - **Nội dung:** {log['content']}")
            
    # 3. Lịch bảo trì hệ thống / Kênh đối tác
    report.append(f"\n## 3. Hoạt động Bảo trì & Giám sát ({len(maintenances)})")
    if not maintenances:
        report.append("- Không có hoạt động bảo trì nào diễn ra trong ca.")
    else:
        for idx, log in enumerate(maintenances, 1):
            report.append(f"{idx}. **{log['component']}** - Trạng thái: [{log['status'].upper()}]")
            report.append(f"   - **Thời gian:** {log['created_at']}")
            report.append(f"   - **Nội dung:** {log['content']}")
            
    # 4. Ghi chú & Việc cần theo dõi tiếp (Action Items)
    report.append(f"\n## 4. Ghi chú ca trực & Công việc cần theo dõi ({len(notes)})")
    if not notes:
        report.append("- Không có ghi chú thêm.")
    else:
        for idx, log in enumerate(notes, 1):
            report.append(f"{idx}. **{log['component']}** - Lực lượng theo dõi: [{log['status'].upper()}]")
            report.append(f"   - **Nội dung:** {log['content']}")
            
    report.append(f"\n---")
    
    sender_name = sender.strip() if (sender and sender.strip()) else "[Họ tên kỹ sư trực]"
    receiver_name = receiver.strip() if (receiver and receiver.strip()) else "[Họ tên kỹ sư ca tiếp theo]"
    report.append(f"\n*Người bàn giao: {sender_name}*")
    report.append(f"\n*Người nhận bàn giao: {receiver_name}*")
    
    return {"markdown": "\n".join(report)}

def parse_nagios_email(subject: str, body: str, date_str: str) -> Dict[str, Any]:
    # Default values
    state = "WARNING"
    host = "Unknown Host"
    service = "Unknown Service"
    message = body.strip().split("\n")[0] if body else subject
    
    # Try parsing subject: ** PROBLEM Service Alert: host1/HTTP is CRITICAL **
    # or ** RECOVERY Service Alert: host1/HTTP is OK **
    # or ** PROBLEM Host Alert: host1 is DOWN **
    subject_lower = subject.lower()
    
    # Check status
    if "critical" in subject_lower or "down" in subject_lower:
        state = "CRITICAL"
    elif "ok" in subject_lower or "recovery" in subject_lower or "up" in subject_lower:
        state = "OK"
    elif "warning" in subject_lower:
        state = "WARNING"
    else:
        state = "WARNING"
        
    # Service Alert matching: host1/HTTP is CRITICAL
    match_service = re.search(r'service alert:\s*([^/]+)/([^\s\*\!]+)\s+is\s+([^\s\*\!]+)', subject, re.IGNORECASE)
    match_host = re.search(r'host alert:\s*([^\s\*\!]+)\s+is\s+([^\s\*\!]+)', subject, re.IGNORECASE)
    
    if match_service:
        host = match_service.group(1).strip()
        service = match_service.group(2).strip()
    elif match_host:
        host = match_host.group(1).strip()
        service = "PING"
    else:
        # Fallback to parsing body
        for line in body.split("\n"):
            line_stripped = line.strip()
            if line_stripped.lower().startswith("host:"):
                host = line_stripped.split(":", 1)[1].strip()
            elif line_stripped.lower().startswith("service:"):
                service = line_stripped.split(":", 1)[1].strip()
                
    # Extract Message (usually "Additional Info:" or similar in body)
    if "Additional Info:" in body:
        parts = body.split("Additional Info:", 1)
        if len(parts) > 1:
            message = parts[1].strip().split("\n")[0].strip()
    elif "info:" in body.lower():
        match_info = re.search(r'info:\s*(.*)', body, re.IGNORECASE)
        if match_info:
            message = match_info.group(1).strip().split("\n")[0].strip()
            
    # Format Date
    try:
        from email.utils import parsedate_to_datetime
        dt = parsedate_to_datetime(date_str)
        formatted_date = dt.strftime("%Y-%m-%d %H:%M:%S")
    except Exception:
        formatted_date = date_str
        
    return {
        "state": state,
        "host": host,
        "service": service,
        "message": message,
        "date": formatted_date,
        "raw_subject": subject
    }

@app.get("/api/nagios-alerts")
def get_nagios_alerts():
    config = read_config()
    imap_enabled = config.get("imap_enabled", False)
    imap_server = config.get("imap_server", "imap.gmail.com")
    imap_user = config.get("imap_user", "")
    imap_pass = config.get("imap_pass", "")
    
    alerts = []
    
    if imap_enabled and imap_server and imap_user and imap_pass:
        import imaplib
        import email
        from email.header import decode_header
        
        try:
            mail = imaplib.IMAP4_SSL(imap_server, port=993)
            mail.login(imap_user, imap_pass)
            mail.select("INBOX")
            
            # Search for emails containing "Nagios"
            status, messages = mail.search(None, '(SUBJECT "Nagios")')
            if status == 'OK':
                mail_ids = messages[0].split()
                # Get latest 10
                mail_ids = mail_ids[-10:]
                mail_ids.reverse()
                
                for mail_id in mail_ids:
                    res, msg_data = mail.fetch(mail_id, '(RFC822)')
                    if res != 'OK':
                        continue
                    for response_part in msg_data:
                        if isinstance(response_part, tuple):
                            msg = email.message_from_bytes(response_part[1])
                            
                            # Decode subject
                            subject, encoding = decode_header(msg["Subject"])[0]
                            if isinstance(subject, bytes):
                                subject = subject.decode(encoding or 'utf-8', errors='ignore')
                                
                            # Decode date
                            date_str = msg["Date"]
                            
                            # Extract body
                            body = ""
                            if msg.is_multipart():
                                for part in msg.walk():
                                    content_type = part.get_content_type()
                                    content_disposition = str(part.get("Content-Disposition"))
                                    if content_type == "text/plain" and "attachment" not in content_disposition:
                                        body_bytes = part.get_payload(decode=True)
                                        body = body_bytes.decode(part.get_content_charset() or 'utf-8', errors='ignore')
                                        break
                            else:
                                body_bytes = msg.get_payload(decode=True)
                                body = body_bytes.decode(msg.get_content_charset() or 'utf-8', errors='ignore')
                                
                            alert = parse_nagios_email(subject, body, date_str)
                            alerts.append(alert)
            mail.logout()
        except Exception as e:
            print(f"Error fetching email from IMAP: {e}")
            pass
            
    # Fallback to Mock Data if no alerts found
    if not alerts:
        now = datetime.datetime.now()
        alerts = [
            {
                "state": "CRITICAL",
                "host": "k8s-prod-node-03",
                "service": "Memory Usage",
                "message": "CRITICAL - Memory usage is 96.5% (Threshold > 95.0%)",
                "date": (now - datetime.timedelta(minutes=5)).strftime("%Y-%m-%d %H:%M:%S"),
                "raw_subject": "** PROBLEM Service Alert: k8s-prod-node-03/Memory Usage is CRITICAL **"
            },
            {
                "state": "WARNING",
                "host": "db-postgres-master",
                "service": "Connection Count",
                "message": "WARNING - Active connections: 452 (Threshold > 400)",
                "date": (now - datetime.timedelta(minutes=14)).strftime("%Y-%m-%d %H:%M:%S"),
                "raw_subject": "** PROBLEM Service Alert: db-postgres-master/Connection Count is WARNING **"
            },
            {
                "state": "CRITICAL",
                "host": "payment-api-gateway",
                "service": "HTTP Response Time",
                "message": "HTTP CRITICAL: 504 Gateway Timeout on /v1/charge",
                "date": (now - datetime.timedelta(minutes=22)).strftime("%Y-%m-%d %H:%M:%S"),
                "raw_subject": "** PROBLEM Service Alert: payment-api-gateway/HTTP Response Time is CRITICAL **"
            },
            {
                "state": "OK",
                "host": "auth-service-02",
                "service": "CPU Load",
                "message": "OK - CPU Load is 12.4% (recovered from CRITICAL)",
                "date": (now - datetime.timedelta(minutes=45)).strftime("%Y-%m-%d %H:%M:%S"),
                "raw_subject": "** RECOVERY Service Alert: auth-service-02/CPU Load is OK **"
            }
        ]
        
    return alerts

@app.post("/api/chat")
def chat_with_agent(req: ChatRequest):
    config = read_config()
    api_key = config.get("api_key", "")
    model_name = config.get("model_name", "minimax/minimax-m2.5")
    
    if not api_key:
        return {"response": "Lỗi: Hiện tại hệ thống chưa cấu hình GreenNode API Key. Vui lòng bấm vào biểu tượng cài đặt bánh răng ở góc phải màn hình để thiết lập API Key."}
        
    logs = read_json_file(LOGS_FILE)
    sops = read_json_file(SOPS_FILE)
    
    # 1. Gather logs summary for injection
    logs_summary = []
    if not logs:
        logs_summary.append("Không có sự cố nào đang ghi nhận.")
    else:
        for log in logs:
            logs_summary.append(
                f"- [Loại: {log['type']}] - Đối tác/Hệ thống: {log['component']} - "
                f"Mức độ: {log['severity']} - Trạng thái: {log['status']} - "
                f"Thời điểm: {log['created_at']} - Nội dung: {log['content']}"
            )
    system_logs_str = "\n".join(logs_summary)
    
    # 2. Simple keyword matching over SOPs (RAG)
    query_text = req.message.lower()
    matched_sops = []
    for sop in sops:
        title_words = set(re_tokenize(sop['title'].lower()))
        content_words = set(re_tokenize(sop['content'].lower()))
        query_words = set(re_tokenize(query_text))
        
        # Simple intersection score
        match_score = len(query_words.intersection(title_words)) * 3 + len(query_words.intersection(content_words))
        if match_score > 0:
            matched_sops.append((match_score, sop))
            
    # Sort by score and get top 3
    matched_sops.sort(key=lambda x: x[0], reverse=True)
    top_sops = [sop for score, sop in matched_sops[:3]]
    
    # If no matches, fall back to top 2 general SOPs just in case
    if not top_sops and sops:
        top_sops = sops[:2]
        
    sops_str = ""
    if not top_sops:
        sops_str = "Hiện chưa có tài liệu quy trình SOP nào được upload lên hệ thống."
    else:
        sops_parts = []
        for sop in top_sops:
            sops_parts.append(f"TÊN QUY TRÌNH: {sop['title']} (File: {sop['filename']})\nNỘI DUNG:\n{sop['content']}\n")
        sops_str = "\n---\n".join(sops_parts)
        
    # 3. Construct System Prompt
    system_prompt = (
        "Bạn là AI Trợ lý Trực ca 247 (247 Operations Duty Assistant) của hệ thống ZaloPay/VNG Cloud.\n"
        "Nhiệm vụ của bạn là hỗ trợ kỹ sư trực ca vận hành hệ thống, tra cứu quy trình vận hành tiêu chuẩn (SOP), soạn thảo email thông báo sự cố (Incident Email) hoặc thông báo bảo trì, và giải đáp thắc mắc.\n\n"
        "Dữ liệu trạng thái hệ thống hiện tại trong ca trực:\n"
        "========================================\n"
        f"{system_logs_str}\n"
        "========================================\n\n"
        "Tài liệu quy trình chuẩn (SOP) đối tác & hệ thống có liên quan:\n"
        "========================================\n"
        f"{sops_str}\n"
        "========================================\n\n"
        "QUY TẮC PHẢN HỒI:\n"
        "1. Trả lời bằng tiếng Việt một cách lịch sự, chuyên nghiệp và ngắn gọn.\n"
        "2. Khi trả lời về quy trình xử lý, hãy tham khảo chính xác từ tài liệu SOP được cung cấp ở trên. Nếu tài liệu không có câu trả lời, hãy báo cho người dùng biết để họ tự upload bổ sung tài liệu (hỗ trợ .docx, .xlsx, .pdf, .txt).\n"
        "3. Khi người dùng yêu cầu soạn thư/email cảnh báo sự cố, hãy tự động trích xuất các thông tin từ 'Dữ liệu trạng thái ca trực' (tên đối tác bị lỗi, thời điểm lỗi, chi tiết lỗi) và điền vào mẫu sau:\n"
        "   - Tiêu đề: [ZaloPay Incident Alert] - Sự cố Kênh [Tên Đối tác] - Mức độ [Severity]\n"
        "   - Nội dung:\n"
        "       + Kính gửi các anh chị,\n"
        "       + Hệ thống ghi nhận sự cố kênh [Tên Đối tác]. Chi tiết sự cố như sau:\n"
        "       + Thời gian phát hiện: [Thời gian]\n"
        "       + Trạng thái: [Trạng thái hiện tại]\n"
        "       + Mức độ ảnh hưởng: [Chi tiết ảnh hưởng]\n"
        "       + Đội vận hành đang phối hợp xử lý và sẽ cập nhật thêm thông tin.\n"
        "4. Tuyệt đối không bịa đặt các thông tin kỹ thuật không có trong tài liệu quy trình."
    )
    
    # 4. Invoke GreenNode MaaS API
    api_url = "https://maas-llm-aiplatform-hcm.api.vngcloud.vn/v1/chat/completions"
    
    # Map chat history to API format
    messages = [{"role": "system", "content": system_prompt}]
    for msg in req.history:
        messages.append({"role": msg.role, "content": msg.content})
        
    messages.append({"role": "user", "content": req.message})
    
    payload = {
        "model": model_name,
        "messages": messages,
        "temperature": 0.2
    }
    
    headers = {
        "Authorization": f"Bearer {api_key}",
        "Content-Type": "application/json"
    }
    
    try:
        response = requests.post(api_url, json=payload, headers=headers, timeout=30)
        if response.status_code == 200:
            res_json = response.json()
            reply = res_json["choices"][0]["message"]["content"]
            return {"response": reply}
        else:
            return {
                "response": f"Lỗi kết nối API GreenNode MaaS (Mã lỗi: {response.status_code}): {response.text}"
            }
    except Exception as e:
        return {"response": f"Gặp lỗi khi gọi AI model: {str(e)}"}

# Helper function to tokenize words for RAG
def re_tokenize(text: str) -> List[str]:
    import re
    # Simple regex word extraction
    return re.findall(r'\w+', text)

# Serve Frontend files
# Serve index.html as root
@app.get("/")
def read_root():
    index_path = os.path.join(FRONTEND_DIR, "index.html")
    if os.path.exists(index_path):
        from fastapi.responses import FileResponse
        return FileResponse(index_path)
    return {"message": "Frontend files not found. Please compile/copy frontend code to frontend/ directory."}

# Mount static files (style.css, app.js, etc.)
app.mount("/", StaticFiles(directory=FRONTEND_DIR), name="static")

if __name__ == "__main__":
    import uvicorn
    uvicorn.run("app:app", host="0.0.0.0", port=8000, reload=True)
